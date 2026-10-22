import streamlit as st
import pandas as pd
import sqlite3
import datetime
import io

# ---------------------------------------------------------
# 1. SETUP PAGE CONFIG & DATABASE STRUCTURE
# ---------------------------------------------------------
st.set_page_config(
    page_title="GS Rurembo Management System",
    page_icon="🏫",
    layout="wide"
)

DB_FILE = "gs_rurembo.db"

def init_db():
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    
    # Table y'Abakoresha Sisitemu (Users & Login)
    c.execute('''
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE,
            password TEXT,
            role TEXT
        )
    ''')
    
    # Ongeramo Admin y'ibanze niba idahari (Username: admin, Password: 123)
    c.execute("SELECT * FROM users WHERE username = 'admin'")
    if not c.fetchone():
        c.execute("INSERT INTO users (username, password, role) VALUES ('admin', '123', 'Administrator')")

    # Table y'abakozi n'Amasezerano
    c.execute('''
        CREATE TABLE IF NOT EXISTS employees (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            emp_id TEXT,
            full_name TEXT,
            phone TEXT,
            position TEXT,
            contract_type TEXT,
            basic_salary REAL,
            start_date TEXT,
            end_date TEXT
        )
    ''')
    
    # Table ya Stock (Ububiko)
    c.execute('''
        CREATE TABLE IF NOT EXISTS stock (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            item_name TEXT UNIQUE,
            category TEXT,
            quantity INTEGER,
            unit TEXT
        )
    ''')
    
    # Table y'Ibyasabwe muri Stock (Stock Requests & Approvals)
    c.execute('''
        CREATE TABLE IF NOT EXISTS stock_requests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            requester_name TEXT,
            item_name TEXT,
            requested_qty INTEGER,
            reason TEXT,
            status TEXT DEFAULT 'PENDING',
            request_date TEXT
        )
    ''')
    
    # Table ya Payroll (Imishahara & RSSB / CBHI)
    c.execute('''
        CREATE TABLE IF NOT EXISTS payroll (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            month_year TEXT,
            full_name TEXT,
            basic_salary REAL,
            allowances REAL,
            pension REAL,
            cbhi REAL,
            maternity REAL,
            other_loans REAL,
            total_deductions REAL,
            net_salary REAL,
            payment_date TEXT
        )
    ''')
    
    for col in ["pension REAL", "cbhi REAL", "maternity REAL", "other_loans REAL", "total_deductions REAL"]:
        try:
            c.execute(f"ALTER TABLE payroll ADD COLUMN {col}")
        except:
            pass

    conn.commit()
    conn.close()

init_db()

# ---------------------------------------------------------
# 2. LOGIN SYSTEM MANAGEMENT
# ---------------------------------------------------------
if 'logged_in' not in st.session_state:
    st.session_state.logged_in = False
    st.session_state.username = ""

def login_screen():
    st.markdown("<h2 style='text-align: center; color: #1E88E5;'>🏫 GS Rurembo Management System</h2>", unsafe_allow_html=True)
    st.markdown("<p style='text-align: center;'>Nyamuneka injira ukoresheje konti yawe.</p>", unsafe_allow_html=True)
    
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        with st.form("login_form"):
            username = st.text_input("Izina ry'ukoresha (Username)")
            password = st.text_input("Ijambo ry'ibanga (Password)", type="password")
            submit = st.form_submit_button("🔓 Injira (Login)")
            
            if submit:
                conn = sqlite3.connect(DB_FILE)
                c = conn.cursor()
                c.execute("SELECT * FROM users WHERE username = ? AND password = ?", (username, password))
                user = c.fetchone()
                conn.close()
                
                if user:
                    st.session_state.logged_in = True
                    st.session_state.username = username
                    st.success("Winjiye neza! Murakaza neza.")
                    st.rerun()
                else:
                    st.error("Izina cyangwa ijambo ry'ibanga sibyo!")

if not st.session_state.logged_in:
    login_screen()
    st.stop()

# ---------------------------------------------------------
# UBURYO BWO KUREMA EXCEL IFITE AMABARA (PROFESSIONAL EXCEL)
# ---------------------------------------------------------
def generate_colored_excel(df, sheet_name="Report"):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        workbook = writer.book
        worksheet = writer.sheets[sheet_name]

        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        data_font = Font(name="Calibri", size=11)
        zebra_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_num in range(2, len(df) + 2):
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.font = data_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")
                if row_num % 2 == 0:
                    cell.fill = zebra_fill

        for col in worksheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            worksheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    return output.getvalue()


# ---------------------------------------------------------
# 3. SIDEBAR NAVIGATION & LOGOUT
# ---------------------------------------------------------
with st.sidebar:
    st.image("https://cdn-icons-png.flaticon.com/512/3135/3135715.png", width=80)
    st.title(f"Winjiye nka: {st.session_state.username}")
    
    if st.button("🔒 Gusohoka (Log out)"):
        st.session_state.logged_in = False
        st.session_state.username = ""
        st.rerun()

    st.markdown("---")
    st.subheader("Hitamo Igice (Menu)")
    
    menu = st.selectbox(
        "Hitamo Igice:",
        [
            "Dashboard",
            "Gucunga Stock / Ububiko",
            "Abakozi n'Amasezerano", 
            "Payroll / Imishahara",
            "Gucunga Users (Admin)",
            "Raporo"
        ]
    )

# ---------------------------------------------------------
# 4. DASHBOARD
# ---------------------------------------------------------
if menu == "Dashboard":
    st.title("🏫 GS Rurembo - Dashboard")
    st.write("Aha ni ku nshamake y'ibikorwa byose n'ububiko muri Sisitemu.")
    st.markdown("---")
    
    conn = sqlite3.connect(DB_FILE)
    df_emp = pd.read_sql_query("SELECT * FROM employees", conn)
    df_stock = pd.read_sql_query("SELECT * FROM stock", conn)
    df_req = pd.read_sql_query("SELECT * FROM stock_requests WHERE status='PENDING'", conn)
    conn.close()

    c1, c2, c3, c4 = st.columns(4)
    with c1:
        st.metric(label="👥 Abakozi Bose", value=len(df_emp))
    with c2:
        st.metric(label="📦 Ubwoko bwa Stock", value=len(df_stock))
    with c3:
        st.metric(label="⏳ Pending Stock Requests", value=len(df_req))
    with c4:
        total_qty = df_stock['quantity'].sum() if not df_stock.empty and 'quantity' in df_stock.columns else 0
        st.metric(label="📊 Ingano y'Ibintu muri Stock", value=int(total_qty))

    st.markdown("---")
    col1, col2 = st.columns(2)
    with col1:
        st.subheader("📦 Ububiko Buri Muri Stock (Stock Balance)")
        if not df_stock.empty:
            st.dataframe(df_stock[['item_name', 'category', 'quantity', 'unit']], use_container_width=True)
        else:
            st.info("Nta gicuruzwa kiragera muri stock.")
    with col2:
        st.subheader("⏳ Ibyasabwe Bitaraprovingwa")
        if not df_req.empty:
            st.dataframe(df_req[['requester_name', 'item_name', 'requested_qty', 'request_date']], use_container_width=True)
        else:
            st.success("Nta busabe bw'ibintu bo muri stock buteregereje!")

# ---------------------------------------------------------
# 5. GUCUNGA STOCK / UBUBIKO
# ---------------------------------------------------------
elif menu == "Gucunga Stock / Ububiko":
    st.title("📦 Icungamutungo ry'Ububiko (Stock Workflow)")
    st.markdown("---")
    
    tab_s1, tab_s2, tab_s3, tab_s4 = st.tabs([
        "📊 Stock Balance", 
        "📥 Kwingiza Item (Stock In)", 
        "📝 Gusaba Item (Request)", 
        "✅ Kwapprova Ibyasabwe (Approvals)"
    ])
    
    with tab_s1:
        st.subheader("📋 Urutonde rw'Ibicuruzwa na Quantity iri muri Ububiko")
        conn = sqlite3.connect(DB_FILE)
        df_stock = pd.read_sql_query("SELECT * FROM stock", conn)
        conn.close()
        
        if df_stock.empty:
            st.info("Nta gicuruzwa kiragera muri stock.")
        else:
            st.dataframe(df_stock, use_container_width=True)
            df_export = df_stock[['item_name', 'category', 'quantity', 'unit']].copy()
            df_export.columns = ["Izina ry'Igicuruzwa", "Icyiciro (Category)", "Ingano (Quantity)", "Igipimo (Unit)"]
            
            excel_data = generate_colored_excel(df_export, sheet_name="Stock_Balance")
            st.download_button(
                label="📥 Kurura Stock Balance (Professional Excel ya Color)",
                data=excel_data,
                file_name=f"GS_Rurembo_Stock_{datetime.date.today()}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

    with tab_s2:
        st.subheader("📥 Ongeraho Igicuruzwa cyangwa Ingano Muri Stock")
        col1, col2 = st.columns(2)
        with col1:
            item_name = st.text_input("Izina ry'Igicuruzwa *")
            category = st.text_input("Icyiciro (Category / Urugero: Isukari, Ibipapuro...)")
        with col2:
            qty = st.number_input("Ingano Uyongereye (Quantity)", min_value=1, step=1)
            unit = st.text_input("Igipimo (e.g. Kg, Pcs, Liters)")
            
        if st.button("💾 Bika / Ongeraho Muri Stock"):
            if item_name:
                try:
                    conn = sqlite3.connect(DB_FILE)
                    c = conn.cursor()
                    c.execute("SELECT quantity FROM stock WHERE item_name = ?", (item_name,))
                    existing = c.fetchone()
                    
                    if existing:
                        new_qty = existing[0] + int(qty)
                        c.execute("UPDATE stock SET quantity = ?, category = ?, unit = ? WHERE item_name = ?", (new_qty, category, unit, item_name))
                        st.success(f"Ingano ya **{item_name}** yavuguruwe! Ingano nshya ni: **{new_qty} {unit}**")
                    else:
                        c.execute("INSERT INTO stock (item_name, category, quantity, unit) VALUES (?, ?, ?, ?)", (item_name, category, int(qty), unit))
                        st.success(f"Igicuruzwa **{item_name}** kiyongereye muri Stock neza!")
                    
                    conn.commit()
                    conn.close()
                    st.rerun()
                except Exception as e:
                    st.error(f"Harimo ikosa: {e}")
            else:
                st.warning("Uzuza izina ry'igicuruzwa!")

    with tab_s3:
        st.subheader("📝 Fomu yo Gusaba Igicuruzwa muri Stock")
        conn = sqlite3.connect(DB_FILE)
        df_stock = pd.read_sql_query("SELECT * FROM stock WHERE quantity > 0", conn)
